"""Generator — 1개 NN.spec.yaml → NN.code.py.

실제 코드 생성은 Claude가 fresh context에서 수행하고, 이 모듈은:
- spec 읽기
- 생성된 코드를 code_path에 저장
- 기본 템플릿(예제) 제공

제공하는 함수:
    generate_stub(spec) -> str   # LLM 없이 최소 동작하는 스켈레톤
    save_code(code, path)

generate-ppt.md v2의 Generator 단계는 Claude가 fresh conversation에서:
1. NN.spec.yaml + template_contract 문서 + ppt_utils 함수 시그니처를 받고
2. NN.code.py 작성
3. Write tool로 저장

후 harness.validator가 검증.
"""

from __future__ import annotations

from pathlib import Path

from harness.schemas import SlideSpec, ProjectPaths


GENERATOR_CONTRACT = """\
당신은 PPT 슬라이드 생성기다. 주어진 SlideSpec 하나만 보고 NN.code.py 를 작성한다.

## 엄격한 제약

1. **import 제약**: `ppt_utils`, `template_contract`, `pptx.*`, stdlib(json/math/re 등)만.
   금지: subprocess, socket, urllib, requests 등.

2. **함수 시그니처 고정**:
   ```python
   def build_slide_N(slide):  # N = spec.idx
       ...
   ```

3. **절대 금지**:
   - prs.slide_masters[*] 에 shape 추가/수정
   - prs.slide_layouts[*] 에 shape 추가/수정
   - 표지 레이아웃 배경 이미지 교체

4. **CONTENT_SAFE 영역**: 모든 custom shape 좌표는 이 영역 안.
   - left ≥ 0.28", top ≥ 0.68"
   - right ≤ 10.56", bottom ≤ 7.02"
   - 이 영역 밖에 shape를 놓으면 로고·푸터를 가려 Validator가 차단.

5. **사용 가능한 헬퍼**: ppt_utils.py 의 함수들만 사용.
   예: set_title, clear_placeholders, add_textbox, add_para, add_rich_text,
       add_accent_bar, make_icon_circle, make_icon_badge, add_shadow,
       add_styled_table, calc_grid, set_body_anchor, add_chevron_step 등.

6. **단일 placeholder**: set_title(slide, "...") 으로 제목을 설정.
   clear_placeholders(slide, keep=[0]) 로 ghost text 제거.
   (layout이 "제목 슬라이드"이면 keep=[0, 1])

7. **색상**: spec.accent_color 사용. 나머지는 자유.

## 기본 스켈레톤

```python
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from ppt_utils import (
    set_title, clear_placeholders, add_textbox, add_para, add_rich_text,
    add_accent_bar, make_icon_circle, make_icon_badge, add_shadow,
    add_styled_table, calc_grid, set_body_anchor, set_shape_opacity,
    add_chevron_step,
)
from template_contract import CONTENT_SAFE, LAYOUT_CONTENT


def build_slide_{IDX}(slide):
    set_title(slide, "{TITLE}")
    clear_placeholders(slide, keep=[0])

    # Key Message Bar
    # ... add_accent_bar + add_textbox at CONTENT_SAFE.top

    # 본문 — spec.content_blocks 에 따라 구성
    # calc_grid(rows, cols) 로 영역 분할 권장

    # CONTENT_SAFE 경계 체크: 모든 shape.left/top/width/height 가 safe_zone 안에 있는지
```

출력: 위 스켈레톤을 채운 완전한 Python 파일 텍스트.
"""


def generate_stub(spec: SlideSpec) -> str:
    """최소 동작 스텁 — LLM 없이 테스트용 코드 생성.

    Validator를 통과하는 합법적 최소 슬라이드를 만든다.
    """
    return f'''"""Auto-generated stub for slide {spec.idx}: {spec.title}."""

from pptx.util import Inches
from ppt_utils import set_title, clear_placeholders, add_textbox


def build_slide_{spec.idx}(slide):
    set_title(slide, {spec.title!r})
    clear_placeholders(slide, keep=[0])
    add_textbox(slide, Inches(1), Inches(2), Inches(8), Inches(3),
                "TODO: generated stub — replace with real content",
                font_size=14)
'''


def save_code(code_text: str, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(code_text, encoding="utf-8")
    return path
