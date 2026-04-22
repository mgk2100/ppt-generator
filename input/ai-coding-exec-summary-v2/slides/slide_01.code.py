"""Slide 1 — AI 코딩 어시스턴트 도입 영향 분석 · 핵심 요약.

Key Message Bar → 3 메시지 카드(1x3) → 승인 요청 박스(2x2 체크박스).
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from ppt_utils import (
    set_title,
    clear_placeholders,
    add_textbox,
    add_para,
    add_rich_text,
    add_accent_bar,
    make_icon_circle,
    set_shape_opacity,
    set_body_anchor,
    calc_grid,
)
from template_contract import CONTENT_SAFE


# --- 팔레트 ------------------------------------------------------------------
C_NAVY  = RGBColor(0x1F, 0x49, 0x7D)
C_GREEN = RGBColor(0x2E, 0x8B, 0x57)
C_BLUE  = RGBColor(0x4F, 0x81, 0xBD)
C_RED   = RGBColor(0xC0, 0x50, 0x4D)
C_INK   = RGBColor(0x1A, 0x1F, 0x2E)
C_GRAY  = RGBColor(0x66, 0x66, 0x66)
C_LGRAY = RGBColor(0xE0, 0xE4, 0xEA)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def build_slide_1(slide):
    set_title(slide, "AI 코딩 어시스턴트 도입 영향 분석  ·  핵심 요약")
    clear_placeholders(slide, keep=[0])

    # ======================================================================
    # 영역 레이아웃 (CONTENT_SAFE: left=0.28, top=0.68, right=10.56, bottom=7.02)
    # - Key Message Bar:   top=0.72,  h=0.52  → bottom=1.24
    # - 3 메시지 카드:       top=1.38,  h=3.50  → bottom=4.88
    # - 승인 요청 박스:      top=5.02,  h=1.95  → bottom=6.97
    # ======================================================================

    cs_left  = CONTENT_SAFE.left
    cs_right = CONTENT_SAFE.right
    cs_width = CONTENT_SAFE.width

    # ---------------------------------------------------------------- 1) Key Message Bar
    kmb_top = Inches(0.72)
    kmb_h   = Inches(0.52)

    # 배경: C_NAVY 10% opacity
    kmb_bg = add_accent_bar(slide, cs_left, kmb_top, cs_width, kmb_h, C_NAVY)
    set_shape_opacity(kmb_bg, 10)

    # 좌측 세로 바 (진한 C_NAVY)
    kmb_side = add_accent_bar(
        slide, cs_left, kmb_top, Inches(0.08), kmb_h, C_NAVY
    )

    # 텍스트 (bold 13pt)
    kmb_tb = add_textbox(
        slide,
        x=cs_left + Inches(0.22), y=kmb_top,
        w=cs_width - Inches(0.32), h=kmb_h,
        text="영역별 차등 적용 시 6개월 내 투자 회수  ·  100명 기준 연 15억 원 절감 가능",
        font_size=13, bold=True, color=C_NAVY, align=PP_ALIGN.LEFT,
    )
    set_body_anchor(kmb_tb, "ctr")

    # ---------------------------------------------------------------- 2) 3 메시지 카드 (1x3)
    card_area = (cs_left, Inches(1.38), cs_width, Inches(3.50))
    grid = calc_grid(1, 3, area=card_area, gap=Inches(0.18))

    cards = [
        {
            "accent": C_GREEN,
            "icon": "✓",
            "title": "경쟁사 이미 도입 중",
            "lines": [
                ("Mercedes",                "5,000명+ 개발자"),
                ("BMW",                     "사내 파일럿 (SPACE 개선)"),
                ("Bosch · 현대모비스 · 현대오토에버", ""),
                ("2027.08",                 "EU AI Act D-Day"),
            ],
        },
        {
            "accent": C_BLUE,
            "icon": "△",
            "title": "영역별 효과 3~5배 차이",
            "lines": [
                ("선행 R&D · 도구",    "+20~35%"),
                ("양산 일반 SW",      "+10~20%"),
                ("양산 안전 코드",    "±5% (본전)"),
                ("→ 선별 적용이 핵심", ""),
            ],
        },
        {
            "accent": C_RED,
            "icon": "⚠",
            "title": "품질검증 부담 증가",
            "lines": [
                ("코드 리뷰 시간",          "+91%"),
                ("코드 리뷰 요청",          "+98%"),
                ("→ 인력·도구 동시 보강 필수", ""),
                ("(가장 자주 누락되는 투자)",  ""),
            ],
        },
    ]

    for idx, cell in enumerate(grid[0]):
        _draw_message_card(slide, cell, cards[idx])

    # ---------------------------------------------------------------- 3) 승인 요청 박스
    _draw_approval_box(slide)


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------

def _draw_message_card(slide, cell, data):
    """단일 메시지 카드: 상단 컬러바 + 원형 아이콘 + 제목 + label:value 리스트."""
    cx, cy, cw, ch = cell.left, cell.top, cell.width, cell.height
    accent = data["accent"]

    # 배경 (아주 옅은 회색 박스)
    bg = add_accent_bar(slide, cx, cy, cw, ch, C_LGRAY)
    set_shape_opacity(bg, 35)

    # 상단 컬러 바
    top_bar_h = Inches(0.08)
    add_accent_bar(slide, cx, cy, cw, top_bar_h, accent)

    # 원형 아이콘 (좌측 상단)
    icon_size = Inches(0.42)
    icon_x = cx + Inches(0.18)
    icon_y = cy + Inches(0.22)
    make_icon_circle(
        slide, icon_x, icon_y, icon_size,
        fill_color=accent, text=data["icon"],
        font_size=12, font_color=C_WHITE,
    )

    # 카드 제목 (아이콘 우측)
    title_x = icon_x + icon_size + Inches(0.10)
    title_w = cx + cw - title_x - Inches(0.12)
    title_tb = add_textbox(
        slide,
        x=title_x, y=icon_y,
        w=title_w, h=icon_size,
        text=data["title"],
        font_size=13, bold=True, color=accent, align=PP_ALIGN.LEFT,
    )
    set_body_anchor(title_tb, "ctr")

    # 얇은 구분선 (제목 아래)
    sep_y = icon_y + icon_size + Inches(0.08)
    add_accent_bar(
        slide,
        cx + Inches(0.18), sep_y,
        cw - Inches(0.36), Emu(9525),
        C_LGRAY,
    )

    # label:value 리스트
    list_x = cx + Inches(0.18)
    list_y = sep_y + Inches(0.12)
    list_w = cw - Inches(0.36)
    list_h = cy + ch - list_y - Inches(0.14)

    list_tb = slide.shapes.add_textbox(int(list_x), int(list_y), int(list_w), int(list_h))
    tf = list_tb.text_frame
    tf.word_wrap = True
    # 첫 paragraph 재사용 (비어있는 기본 paragraph)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT

    for i, (label, value) in enumerate(data["lines"]):
        segments = []
        if label:
            lbl_bold = label.startswith("→") or label.startswith("(")
            lbl_color = accent if label.startswith("→") else (C_GRAY if label.startswith("(") else C_INK)
            segments.append({
                "text": label,
                "font_size": 10,
                "bold": lbl_bold,
                "color": lbl_color,
            })
        if value:
            segments.append({"text": "  ", "font_size": 10})
            segments.append({
                "text": value,
                "font_size": 10,
                "bold": True,
                "color": accent,
            })

        if i == 0:
            # 첫 paragraph를 이용: 직접 runs 채우기
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            for seg in segments:
                run = p.add_run()
                run.text = str(seg["text"])
                run.font.size = Pt(seg.get("font_size", 10))
                if seg.get("bold"):
                    run.font.bold = True
                if "color" in seg:
                    run.font.color.rgb = seg["color"]
        else:
            add_rich_text(
                tf, segments,
                align=PP_ALIGN.LEFT,
                space_before=Pt(4),
                line_spacing=Pt(14),
            )


def _draw_approval_box(slide):
    """하단 승인 요청 박스: C_NAVY 8% 배경 + 제목 + 2x2 체크박스 항목."""
    cs_left  = CONTENT_SAFE.left
    cs_width = CONTENT_SAFE.width

    box_top = Inches(5.02)
    box_h   = Inches(1.95)

    # 배경 (C_NAVY 8% opacity)
    bg = add_accent_bar(slide, cs_left, box_top, cs_width, box_h, C_NAVY)
    set_shape_opacity(bg, 8)

    # 좌측 세로 바
    add_accent_bar(slide, cs_left, box_top, Inches(0.08), box_h, C_NAVY)

    # 제목
    title_tb = add_textbox(
        slide,
        x=cs_left + Inches(0.22), y=box_top + Inches(0.10),
        w=cs_width - Inches(0.32), h=Inches(0.40),
        text="▶ 즉시 결정 요청 사항",
        font_size=13, bold=True, color=C_NAVY, align=PP_ALIGN.LEFT,
    )
    set_body_anchor(title_tb, "ctr")

    # 2x2 체크박스 항목
    items = [
        "선행 · 테스트 자동화 영역  즉시 도입 승인",
        "4개월 시범 사업 예산 승인 (약 3천만원)",
        "품질검증 인력 보강 계획  착수",
        "사내 AI 사용 가이드라인  제정",
    ]

    items_area = (
        cs_left + Inches(0.30),
        box_top + Inches(0.58),
        cs_width - Inches(0.50),
        Inches(1.25),
    )
    cells = calc_grid(2, 2, area=items_area, gap=Inches(0.14))

    for i, item in enumerate(items):
        r, c = divmod(i, 2)
        cell = cells[r][c]
        _draw_check_item(slide, cell, item)


def _draw_check_item(slide, cell, text):
    """체크박스 □ + 항목 텍스트."""
    cb_size = Inches(0.22)
    cb_x = cell.left
    cb_y = cell.top + (cell.height - cb_size) // 2

    # 체크박스 (빈 사각형, 진한 테두리)
    cb = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, cb_x, cb_y, cb_size, cb_size
    )
    cb.fill.solid()
    cb.fill.fore_color.rgb = C_WHITE
    cb.line.color.rgb = C_NAVY
    cb.line.width = Pt(1.25)

    # 항목 텍스트
    txt_x = cb_x + cb_size + Inches(0.12)
    txt_w = cell.left + cell.width - txt_x
    tb = add_textbox(
        slide,
        x=txt_x, y=cell.top,
        w=txt_w, h=cell.height,
        text=text,
        font_size=11, bold=False, color=C_INK, align=PP_ALIGN.LEFT,
    )
    set_body_anchor(tb, "ctr")
