"""Slide 03 — 경쟁사 상세 (도구 · 규모 · 측정)."""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from ppt_utils import (
    set_title, clear_placeholders,
    add_textbox, add_rich_text,
    add_accent_bar, make_icon_badge,
    add_shadow, set_body_anchor, set_text_inset,
    calc_grid,
)
from template_contract import CONTENT_SAFE


# ---- 색상 ------------------------------------------------------------
C_NAVY  = RGBColor(0x1F, 0x49, 0x7D)
C_BLUE  = RGBColor(0x4F, 0x81, 0xBD)
C_TEAL  = RGBColor(0x2C, 0x7F, 0x94)
C_INK   = RGBColor(0x1A, 0x1F, 0x2E)
C_GRAY  = RGBColor(0x66, 0x66, 0x66)
C_LGRAY = RGBColor(0xE0, 0xE4, 0xEA)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ---- 카드 데이터 -----------------------------------------------------
CARDS = [
    {
        "name": "Mercedes-Benz",
        "badge": "2023.07~",
        "badge_color": C_NAVY,
        "tool":        "GitHub Copilot",
        "scale":       "5,000명+ · 115k repos · 4,300 Enterprise orgs",
        "measurement": "개발자 자체 설문 + 흐름 상태 보고",
        "result":      "주당 30분+ 절감 · 누적 200만 라인 수락",
        "source":      "GitHub 공식 case study",
    },
    {
        "name": "BMW Group",
        "badge": "2024",
        "badge_color": C_NAVY,
        "tool":        "GitHub Copilot (사내 파일럿)",
        "scale":       "사내 파일럿 (규모 비공개)",
        "measurement": "SPACE 프레임워크 5축 (S·P·A·C·E)",
        "result":      "5축 전항목 개선 · 결함 감소 (정량치 비공개)",
        "source":      "AMCIS 2024 논문 (Pielmeier·Eidelloth)",
    },
    {
        "name": "Bosch",
        "badge": "진행 중",
        "badge_color": C_NAVY,
        "tool":        "GitHub Copilot (bosch-copilot org)",
        "scale":       "사내 org 운영 (액세스 관리)",
        "measurement": "비공개",
        "result":      "단계적 확산 중",
        "source":      "GitHub 공식 org 페이지",
    },
    {
        "name": "현대모비스",
        "badge": "2025.09",
        "badge_color": C_TEAL,
        "tool":        "Mobis Development Studio (Wind River 협업)",
        "scale":       "웹 기반 통합 SDV 개발환경",
        "measurement": "CI/CD/CT 자동화 지표",
        "result":      "차세대 개발시스템 확장 · shift-left 테스팅",
        "source":      "Wind River 보도자료 (2025.09)",
    },
    {
        "name": "현대오토에버",
        "badge": "2024~",
        "badge_color": C_TEAL,
        "tool":        "H-Chat (Azure OpenAI / Gemini / Claude 프록시)",
        "scale":       "그룹사 전사 배포",
        "measurement": "비공개 (정량 미공개)",
        "result":      "코드 보조 · 문서 작성 지원",
        "source":      "현대오토에버 공식 페이지",
    },
    {
        "name": "Geely Auto",
        "badge": "2025.07",
        "badge_color": C_TEAL,
        "tool":        "AI 안전 프로세스 (코딩 도구 아님)",
        "scale":       "차량 기능 AI 전반",
        "measurement": "ISO/PAS 8800 인증 심사 (SGS-TÜV Saar)",
        "result":      "세계 최초 AI 안전 인증 취득",
        "source":      "SGS 보도자료 (2025.08)",
    },
]


# ---- 빌더 -----------------------------------------------------------


def _build_key_message_bar(slide, x, y, w, h, text):
    """상단 Key Message Bar (좌측 accent 세로바 + 본문 텍스트 배경)."""
    # 배경 bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0xF2, 0xF5, 0xFA)
    bar.line.color.rgb = C_LGRAY
    bar.line.width = Pt(0.5)

    # 좌측 accent 세로바
    add_accent_bar(slide, x, y, Inches(0.08), h, C_BLUE)

    # 텍스트
    tb = add_textbox(
        slide,
        x + Inches(0.20), y, w - Inches(0.24), h,
        text,
        font_size=12, color=C_INK, bold=True,
        align=PP_ALIGN.LEFT,
    )
    set_body_anchor(tb, "ctr")
    set_text_inset(tb, left=Inches(0.10), right=Inches(0.08),
                   top=Inches(0.04), bottom=Inches(0.04))
    return bar


def _build_card(slide, rect, data):
    """단일 경쟁사 상세 카드. 5행 구조."""
    x, y, w, h = rect.left, rect.top, rect.width, rect.height

    # 카드 배경
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_LGRAY
    card.line.width = Pt(0.75)
    add_shadow(card, blur_pt=4, dist_pt=2, opacity_pct=18)

    # 상단 accent 가로바
    accent_h = Inches(0.05)
    add_accent_bar(slide, x, y, w, accent_h, C_BLUE)

    # ---- 헤더 (회사명 + 시점 배지) ----
    header_y = y + accent_h + Inches(0.08)
    header_h = Inches(0.34)

    # 배지 (우측). 폭은 텍스트 길이에 따라 조정 (8자 기준)
    badge_text = data["badge"]
    badge_w = Inches(1.05) if len(badge_text) >= 8 else Inches(0.90)
    badge_h = Inches(0.28)
    badge_x = x + w - badge_w - Inches(0.12)
    badge_y = header_y + (header_h - badge_h) // 2
    make_icon_badge(
        slide, badge_x, badge_y, badge_w, badge_h,
        badge_text, data["badge_color"],
        font_size=9, font_color=C_WHITE, corner_radius=0.3,
    )

    # 회사명 (좌측). 배지 영역 침범 방지
    name_w = badge_x - x - Inches(0.24)
    name_tb = add_textbox(
        slide, x + Inches(0.14), header_y, name_w, header_h,
        data["name"],
        font_size=12, color=C_INK, bold=True, align=PP_ALIGN.LEFT,
    )
    set_body_anchor(name_tb, "ctr")
    set_text_inset(name_tb, left=Inches(0.02), right=Inches(0.02),
                   top=Inches(0.02), bottom=Inches(0.02))

    # ---- Divider ----
    div_y = header_y + header_h + Inches(0.04)
    add_accent_bar(
        slide,
        x + Inches(0.14), div_y,
        w - Inches(0.28), Inches(0.012),
        C_LGRAY,
    )

    # ---- 5행 상세 (단일 textbox에 rich-text 5개 paragraph) ----
    rows = [
        ("도구", data["tool"]),
        ("규모", data["scale"]),
        ("측정", data["measurement"]),
        ("결과", data["result"]),
        ("출처", data["source"]),
    ]

    body_y_start = div_y + Inches(0.10)
    body_bottom  = y + h - Inches(0.08)
    body_h_total = body_bottom - body_y_start
    body_left = x + Inches(0.14)
    body_w = w - Inches(0.28)

    body_tb = slide.shapes.add_textbox(
        body_left, body_y_start, body_w, body_h_total,
    )
    tf = body_tb.text_frame
    tf.word_wrap = True
    # 기본 paragraph는 비워두고 add_rich_text가 add_paragraph로 추가하도록 하되,
    # 첫 paragraph가 빈 채로 남지 않도록 텍스트 설정 전에 초기 단락 사용.
    # add_rich_text는 항상 add_paragraph()를 호출 → 첫 빈 단락 존재 → 제거 필요.
    set_text_inset(body_tb, left=Inches(0.02), right=Inches(0.02),
                   top=Inches(0.02), bottom=Inches(0.02))

    for i, (label, value) in enumerate(rows):
        add_rich_text(
            tf,
            [
                {"text": f"{label}  ", "font_size": 10, "bold": True, "color": C_BLUE},
                {"text": value,        "font_size": 10, "color": C_INK},
            ],
            align=PP_ALIGN.LEFT,
            space_before=Pt(1) if i > 0 else Pt(0),
            space_after=Pt(1),
            line_spacing=Pt(12),
        )

    # add_rich_text가 매번 add_paragraph하므로 초기 빈 단락 제거
    _tb = tf._txBody
    from pptx.oxml.ns import qn as _qn
    paras = _tb.findall(_qn("a:p"))
    if paras and (paras[0].find(_qn("a:r")) is None and paras[0].find(_qn("a:fld")) is None):
        _tb.remove(paras[0])


def build_slide_3(slide):
    set_title(slide, "경쟁사 상세  ·  도구 · 규모 · 측정")
    clear_placeholders(slide, keep=[0])

    # ---- Key Message Bar ----
    km_x = CONTENT_SAFE.left
    km_y = CONTENT_SAFE.top
    km_w = CONTENT_SAFE.width
    km_h = Inches(0.50)
    _build_key_message_bar(
        slide, km_x, km_y, km_w, km_h,
        "주요 OEM·Tier-1 6사 도입 현황 (공식 공표 기준)  ·  "
        "2027.08 EU AI Act 자동차 D-Day",
    )

    # ---- 카드 영역 (2x3 그리드) ----
    grid_top = km_y + km_h + Inches(0.14)
    grid_bottom = CONTENT_SAFE.bottom - Inches(0.02)
    grid_area = (
        CONTENT_SAFE.left,
        grid_top,
        CONTENT_SAFE.width,
        grid_bottom - grid_top,
    )
    grid = calc_grid(2, 3, area=grid_area, gap=Inches(0.15))

    for i, data in enumerate(CARDS):
        r, c = divmod(i, 3)
        _build_card(slide, grid[r][c], data)
