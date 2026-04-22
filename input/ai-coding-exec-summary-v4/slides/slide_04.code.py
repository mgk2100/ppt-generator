"""Slide 4 — 영역별 예상 효과 · 외부 공식 벤치마크."""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from ppt_utils import (
    set_title, clear_placeholders,
    add_textbox, add_para, add_rich_text,
    add_accent_bar,
    set_body_anchor, set_text_inset,
)
from template_contract import CONTENT_SAFE, LAYOUT_CONTENT


# ============================ 색상 팔레트 ============================
C_NAVY   = RGBColor(0x1F, 0x49, 0x7D)
C_BLUE   = RGBColor(0x4F, 0x81, 0xBD)
C_GREEN  = RGBColor(0x2E, 0x8B, 0x57)
C_YELLOW = RGBColor(0xD4, 0xA0, 0x17)
C_RED    = RGBColor(0xC0, 0x50, 0x4D)
C_INK    = RGBColor(0x1A, 0x1F, 0x2E)
C_GRAY   = RGBColor(0x66, 0x66, 0x66)
C_LGRAY  = RGBColor(0xE0, 0xE4, 0xEA)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BG_GREEN  = RGBColor(0xE8, 0xF5, 0xEC)
BG_YELLOW = RGBColor(0xFF, 0xF7, 0xD9)
BG_RED    = RGBColor(0xFC, 0xE8, 0xE6)
BG_KEY    = RGBColor(0xEE, 0xF3, 0xFA)


def _add_filled_rect(slide, x, y, w, h, fill, line_color=None, line_pt=0.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_pt)
    return shape


def _add_dot(slide, cx, cy, diameter, color):
    d = int(diameter)
    left = int(cx) - d // 2
    top = int(cy) - d // 2
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, d, d)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _clean_leading_empty(tf):
    """add_rich_text 가 남기는 빈 첫 단락 제거."""
    if len(tf.paragraphs) > 1 and not tf.paragraphs[0].text:
        p0 = tf.paragraphs[0]
        p0._p.getparent().remove(p0._p)


def build_slide_4(slide):
    set_title(slide, "영역별 예상 효과  ·  외부 공식 벤치마크")
    clear_placeholders(slide, keep=[0])

    SL = CONTENT_SAFE.left
    ST = CONTENT_SAFE.top
    SW = CONTENT_SAFE.width
    SB = CONTENT_SAFE.bottom

    # ============================================================
    # 1) Key Message Bar
    # ============================================================
    km_h = Inches(0.46)
    km_y = ST
    # accent 좌측 바 + 배경 rect (bg는 accent와 나란히)
    add_accent_bar(slide, SL, km_y, Inches(0.08), km_h, C_NAVY)
    _add_filled_rect(
        slide,
        SL + Inches(0.08), km_y,
        SW - Inches(0.08), km_h,
        BG_KEY,
    )
    km_tb = add_textbox(
        slide,
        SL + Inches(0.22), km_y,
        SW - Inches(0.30), km_h,
        "Key Message   |   영역별 효과 3~5배 차이 · 선행·도구 영역이 즉시 도입 후보",
        font_size=13, bold=True, color=C_NAVY,
        align=PP_ALIGN.LEFT,
    )
    set_body_anchor(km_tb, 'ctr')
    set_text_inset(km_tb, left=Inches(0.10), right=Inches(0.10),
                   top=Inches(0.04), bottom=Inches(0.04))

    # ============================================================
    # 2) ① 신호등 테이블
    # ============================================================
    sec1_y = km_y + km_h + Inches(0.12)
    sec1_title_h = Inches(0.26)
    add_textbox(
        slide,
        SL, sec1_y, SW, sec1_title_h,
        "① 영역별 효과 (업계 일반 보정 추정)",
        font_size=11, bold=True, color=C_INK,
        align=PP_ALIGN.LEFT,
    )

    tbl_y = sec1_y + sec1_title_h + Inches(0.02)
    header_h = Inches(0.30)
    row_h = Inches(0.38)
    tbl_h = header_h + row_h * 4
    tbl_w = SW

    # 컬럼 비율: 영역 : 효과 : 해석
    col_ratio = [45, 22, 33]
    total_ratio = sum(col_ratio)
    col_w = [int(tbl_w) * r // total_ratio for r in col_ratio]
    col_w[-1] = int(tbl_w) - sum(col_w[:-1])
    col_x = [int(SL)]
    col_x.append(col_x[0] + col_w[0])
    col_x.append(col_x[1] + col_w[1])

    # 헤더: 단일 bg rect + 3 label textboxes
    _add_filled_rect(
        slide,
        int(SL), tbl_y, int(tbl_w), header_h,
        C_NAVY,
    )
    header_names = ["영역", "효과", "해석"]
    header_aligns = [PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.LEFT]
    for i, name in enumerate(header_names):
        htb = add_textbox(
            slide,
            col_x[i], tbl_y,
            col_w[i], header_h,
            name,
            font_size=10, bold=True, color=C_WHITE,
            align=header_aligns[i],
        )
        set_body_anchor(htb, 'ctr')
        set_text_inset(htb, left=Inches(0.16), right=Inches(0.10),
                       top=Inches(0.02), bottom=Inches(0.02))

    # 4개 행 데이터
    rows = [
        {"dot": C_GREEN,  "bg": BG_GREEN,
         "area": "선행 R&D · 도구 · 테스트 자동화",
         "eff": "+20 ~ +35%",   "eff_col": C_GREEN,
         "inter": "즉시 도입 후보 (위험 낮음 · 효과 큼)"},
        {"dot": C_GREEN,  "bg": BG_GREEN,
         "area": "양산 일반 SW (비안전)",
         "eff": "+10 ~ +20%",   "eff_col": C_GREEN,
         "inter": "코드 검사 도구 통합 후 도입"},
        {"dot": C_YELLOW, "bg": BG_YELLOW,
         "area": "양산 안전 코드 (차량 안전등급 높음↑)",
         "eff": "-5 ~ +10%",    "eff_col": C_YELLOW,
         "inter": "본전 영역 · 자사 시범 측정 필요"},
        {"dot": C_RED,    "bg": BG_RED,
         "area": "품질검증 단계 (인력 관점)",
         "eff": "-15 ~ -30%",   "eff_col": C_RED,
         "inter": "코드 리뷰 부담 증가 · 인력 보강 필수"},
    ]

    ry = tbl_y + header_h
    for r in rows:
        # 행 배경
        _add_filled_rect(
            slide,
            int(SL), ry,
            int(tbl_w), row_h,
            r["bg"],
        )

        # OVAL dot (영역 칼럼 좌측)
        dot_d = Inches(0.14)
        dot_cx = col_x[0] + Inches(0.18)
        dot_cy = ry + row_h // 2
        _add_dot(slide, dot_cx, dot_cy, dot_d, r["dot"])

        # 영역 텍스트
        area_tb = add_textbox(
            slide,
            col_x[0] + Inches(0.34), ry,
            col_w[0] - Inches(0.40), row_h,
            r["area"],
            font_size=10.5, bold=False, color=C_INK,
            align=PP_ALIGN.LEFT,
        )
        set_body_anchor(area_tb, 'ctr')
        set_text_inset(area_tb, left=Inches(0.04), right=Inches(0.04),
                       top=Inches(0.02), bottom=Inches(0.02))

        # 효과 (bold accent center)
        eff_tb = add_textbox(
            slide,
            col_x[1], ry,
            col_w[1], row_h,
            r["eff"],
            font_size=14, bold=True, color=r["eff_col"],
            align=PP_ALIGN.CENTER,
        )
        set_body_anchor(eff_tb, 'ctr')
        set_text_inset(eff_tb, left=Inches(0.04), right=Inches(0.04),
                       top=Inches(0.02), bottom=Inches(0.02))

        # 해석
        inter_tb = add_textbox(
            slide,
            col_x[2], ry,
            col_w[2], row_h,
            r["inter"],
            font_size=10, bold=False, color=C_INK,
            align=PP_ALIGN.LEFT,
        )
        set_body_anchor(inter_tb, 'ctr')
        set_text_inset(inter_tb, left=Inches(0.10), right=Inches(0.08),
                       top=Inches(0.02), bottom=Inches(0.02))

        ry += row_h

    # ============================================================
    # 3) ② 벤치마크 상세 카드 (1x3)
    # ============================================================
    sec2_y = tbl_y + tbl_h + Inches(0.14)
    sec2_title_h = Inches(0.24)
    add_textbox(
        slide,
        SL, sec2_y, SW, sec2_title_h,
        "② 외부 공식 벤치마크 (상세)",
        font_size=11, bold=True, color=C_INK,
        align=PP_ALIGN.LEFT,
    )

    cards_y_top = sec2_y + sec2_title_h + Inches(0.04)
    cards_y_bot = SB - Inches(0.02)
    cards_h     = cards_y_bot - cards_y_top

    cards = [
        {"source": "Mercedes-Benz",
         "source_sub": "GitHub 공식 case study (2023.07~)",
         "tool": "GitHub Copilot",
         "scale": "5,000명+ 개발자 · 115k repos",
         "meas": "개발자 자체 설문",
         "headline": "주당 30분+ 절감",
         "detail": "누적 200만 라인 수락 · 흐름 상태 유지 향상",
         "accent": C_NAVY},
        {"source": "BMW Group",
         "source_sub": "AMCIS 2024 논문",
         "tool": "GitHub Copilot (사내 파일럿)",
         "scale": "사내 파일럿",
         "meas": "SPACE 프레임워크 5축",
         "headline": "5축 전항목 개선",
         "detail": "결함 감소 (정량치 비공개) · Pielmeier·Eidelloth",
         "accent": C_NAVY},
        {"source": "Faros AI (2025)",
         "source_sub": "다수 엔터프라이즈 분석",
         "tool": "Copilot / Cursor 등 혼합",
         "scale": "다수 엔터프라이즈 (규모 비공개)",
         "meas": "Git 로그 · 리뷰 메트릭 분석",
         "headline": "리뷰 시간 +91% · 요청 +98%",
         "detail": "작업 처리량 +21% · AI 코드량 증가가 검증 부담으로 직결",
         "accent": C_BLUE},
    ]

    card_gap = Inches(0.12)
    total_gap = int(card_gap) * (len(cards) - 1)
    card_w = (int(SW) - total_gap) // len(cards)

    pad_x = Inches(0.14)

    for i, c in enumerate(cards):
        cx = int(SL) + i * (card_w + int(card_gap))
        cy = int(cards_y_top)
        ch = int(cards_h)

        # 카드 배경 + 테두리
        _add_filled_rect(
            slide,
            cx, cy, card_w, ch,
            C_WHITE,
            line_color=C_LGRAY, line_pt=0.75,
        )
        # 상단 accent 바
        top_bar_h = Inches(0.10)
        _add_filled_rect(
            slide,
            cx, cy, card_w, top_bar_h,
            c["accent"],
        )

        inner_x = cx + int(pad_x)
        inner_w = card_w - 2 * int(pad_x)

        # 상단 영역: 출처 + 공개매체 + divider 를 하나의 header_tb 로 — 단 divider 는 rect 로
        header_y = cy + int(top_bar_h) + int(Inches(0.10))
        header_h2 = Inches(0.54)
        header_tb = slide.shapes.add_textbox(inner_x, header_y, inner_w, int(header_h2))
        htf = header_tb.text_frame
        htf.word_wrap = True
        add_rich_text(
            htf,
            [{"text": c["source"], "font_size": 13, "color": C_INK, "bold": True}],
            align=PP_ALIGN.LEFT,
        )
        add_rich_text(
            htf,
            [{"text": c["source_sub"], "font_size": 9, "color": C_GRAY, "bold": False}],
            align=PP_ALIGN.LEFT, space_before=Pt(2),
        )
        _clean_leading_empty(htf)
        set_text_inset(header_tb, left=Inches(0.02), right=Inches(0.02),
                       top=Inches(0.0), bottom=Inches(0.0))

        div_y = header_y + int(header_h2) + int(Inches(0.04))
        div_h = Inches(0.012)
        _add_filled_rect(
            slide,
            inner_x, div_y, inner_w, div_h,
            C_LGRAY,
        )

        # 본문: 메타 3줄 + 핵심 수치 + 상세 를 하나의 textbox 로 합침
        body_y = div_y + int(div_h) + int(Inches(0.08))
        body_bottom = cy + ch - int(Inches(0.10))
        body_h = max(int(Inches(1.8)), body_bottom - body_y)

        body_tb = slide.shapes.add_textbox(inner_x, body_y, inner_w, body_h)
        btf = body_tb.text_frame
        btf.word_wrap = True

        # 도구 / 규모 / 측정 (label 9pt gray bold + value 10pt ink)
        for label, value in [("도구", c["tool"]),
                             ("규모", c["scale"]),
                             ("측정", c["meas"])]:
            add_rich_text(
                btf,
                [
                    {"text": f"{label}  ", "font_size": 9, "color": C_GRAY, "bold": True},
                    {"text": value, "font_size": 10, "color": C_INK, "bold": False},
                ],
                align=PP_ALIGN.LEFT, space_before=Pt(2), space_after=Pt(2),
                line_spacing=Pt(13),
            )

        # 핵심 수치 (★ 14pt bold accent)
        add_rich_text(
            btf,
            [
                {"text": "\u2605 ", "font_size": 14, "color": c["accent"], "bold": True},
                {"text": c["headline"], "font_size": 14, "color": c["accent"], "bold": True},
            ],
            align=PP_ALIGN.LEFT, space_before=Pt(8), space_after=Pt(0),
            line_spacing=Pt(17),
        )

        # 상세 (└ 9pt gray)
        add_rich_text(
            btf,
            [
                {"text": "\u2514 ", "font_size": 9, "color": C_GRAY, "bold": False},
                {"text": c["detail"], "font_size": 9, "color": C_GRAY, "bold": False},
            ],
            align=PP_ALIGN.LEFT, space_before=Pt(2), space_after=Pt(0),
            line_spacing=Pt(12),
        )

        _clean_leading_empty(btf)
        set_text_inset(body_tb, left=Inches(0.02), right=Inches(0.02),
                       top=Inches(0.0), bottom=Inches(0.02))
