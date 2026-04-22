from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from ppt_utils import (
    set_title, clear_placeholders,
    add_textbox, add_para, add_rich_text,
    add_accent_bar, make_icon_badge,
    add_shadow, set_shape_opacity,
    set_body_anchor, set_text_inset,
    calc_grid,
)
from template_contract import CONTENT_SAFE, LAYOUT_CONTENT


# ---------- 색상 ----------
C_NAVY  = RGBColor(0x1F, 0x49, 0x7D)
C_GREEN = RGBColor(0x2E, 0x8B, 0x57)
C_BLUE  = RGBColor(0x4F, 0x81, 0xBD)
C_RED   = RGBColor(0xC0, 0x50, 0x4D)
C_INK   = RGBColor(0x1A, 0x1F, 0x2E)
C_GRAY  = RGBColor(0x66, 0x66, 0x66)
C_LGRAY = RGBColor(0xE0, 0xE4, 0xEA)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def build_slide_2(slide):
    set_title(slide, "핵심 요약")
    clear_placeholders(slide, keep=[0])

    # ============================================================
    # 본문 영역 정의: y=0.68"~6.55" (h=5.87")
    # 각주 영역 정의: y=6.60"~7.00" (h=0.40")
    # ============================================================
    body_left = CONTENT_SAFE.left
    body_width = CONTENT_SAFE.width
    body_top = Inches(0.68)
    body_bottom = Inches(6.55)
    body_height = body_bottom - body_top  # 5.87"

    # ============================================================
    # 1) Key Message Bar (상단) — accent 바 + 메시지 텍스트
    # ============================================================
    kmb_y = body_top
    kmb_h = Inches(0.55)

    # 좌측 accent 바 (navy)
    accent_bar_w = Inches(0.08)
    add_accent_bar(slide, body_left, kmb_y, accent_bar_w, kmb_h, C_NAVY)

    # Key Message 배경 패널 (아주 옅은 회색)
    km_panel_left = body_left + accent_bar_w
    km_panel_w = body_width - accent_bar_w
    km_panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, km_panel_left, kmb_y, km_panel_w, kmb_h
    )
    km_panel.fill.solid()
    km_panel.fill.fore_color.rgb = C_LGRAY
    set_shape_opacity(km_panel, 35)
    km_panel.line.fill.background()

    # Key Message 텍스트 (rich)
    km_tb = slide.shapes.add_textbox(
        km_panel_left, kmb_y, km_panel_w, kmb_h
    )
    km_tf = km_tb.text_frame
    km_tf.word_wrap = True
    set_text_inset(km_tb, left=Inches(0.18), right=Inches(0.18),
                   top=Inches(0.08), bottom=Inches(0.08))
    set_body_anchor(km_tb, 'ctr')

    # 첫 단락을 지우고 rich로 넣기
    km_tf.paragraphs[0].text = ""
    add_rich_text(km_tf, [
        {"text": "경쟁사는 이미 ", "font_size": 13, "color": C_INK, "bold": True},
        {"text": "도입 단계", "font_size": 13, "color": C_NAVY, "bold": True},
        {"text": "  ·  영역별 효과 차이 크므로 ", "font_size": 13, "color": C_INK, "bold": True},
        {"text": "선별 적용 필수", "font_size": 13, "color": C_NAVY, "bold": True},
    ], align=PP_ALIGN.LEFT)

    # ============================================================
    # 2) 3 메시지 카드 (1x3 grid) — Key Message 아래 ~ takeaway 위
    # ============================================================
    cards_top = kmb_y + kmb_h + Inches(0.18)
    # takeaway bar 높이 + 상단 간격 예약
    takeaway_h = Inches(1.30)
    takeaway_gap_top = Inches(0.18)
    cards_bottom = body_bottom - takeaway_h - takeaway_gap_top
    cards_h = cards_bottom - cards_top

    grid = calc_grid(
        1, 3,
        area=(body_left, cards_top, body_width, cards_h),
        gap=Inches(0.20),
    )

    cards_data = [
        {
            "accent": C_GREEN,
            "icon": "✓",
            "title": "경쟁사 이미 도입 중",
            "lines": [
                ("Mercedes", "5,000명+ Copilot 전사 사용"),
                ("BMW", "사내 파일럿 · SPACE 5축 개선"),
                ("현대모비스", "Mobis Development Studio"),
                ("현대오토에버", "H-Chat 그룹사 전사"),
            ],
        },
        {
            "accent": C_BLUE,
            "icon": "△",
            "title": "영역별 효과 차이 (업계 일반)",
            "lines": [
                ("선행 R&D · 도구", "+20~35%"),
                ("양산 일반 SW", "+10~20%"),
                ("양산 안전 코드", "±5% (본전)"),
                ("→ 선별 적용이 핵심", ""),
            ],
        },
        {
            "accent": C_RED,
            "icon": "⚠",
            "title": "품질검증 부담 증가",
            "lines": [
                ("코드 리뷰 시간", "+91% (Faros 2025)"),
                ("코드 리뷰 요청", "+98% (Faros 2025)"),
                ("→ 인력·도구 동시 보강 필수", ""),
                ("(가장 자주 누락되는 투자)", ""),
            ],
        },
    ]

    for idx, cell in enumerate(grid[0]):
        card = cards_data[idx]

        # 카드 컨테이너 (흰 바탕 + 옅은 테두리)
        card_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, cell.left, cell.top, cell.width, cell.height
        )
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = C_WHITE
        card_shape.line.color.rgb = C_LGRAY
        card_shape.line.width = Pt(0.75)
        add_shadow(card_shape, blur_pt=4, dist_pt=2,
                   opacity_pct=15, color=C_GRAY)

        # 상단 accent 바
        top_bar_h = Inches(0.06)
        add_accent_bar(slide, cell.left, cell.top,
                       cell.width, top_bar_h, card["accent"])

        # 헤더 영역 (아이콘 배지 + 타이틀)
        header_y = cell.top + top_bar_h + Inches(0.10)
        header_h = Inches(0.42)

        badge_size = Inches(0.36)
        badge_x = cell.left + Inches(0.15)
        badge_y = header_y + (header_h - badge_size) // 2
        make_icon_badge(
            slide, badge_x, badge_y, badge_size, badge_size,
            text=card["icon"], fill_color=card["accent"],
            font_size=14, font_color=C_WHITE,
            corner_radius=0.30,
        )

        # 타이틀
        title_x = badge_x + badge_size + Inches(0.10)
        title_w = cell.width - (title_x - cell.left) - Inches(0.10)
        title_tb = add_textbox(
            slide, title_x, header_y, title_w, header_h,
            card["title"], font_size=12, bold=True, color=C_INK,
            align=PP_ALIGN.LEFT,
        )
        set_body_anchor(title_tb, 'ctr')
        set_text_inset(title_tb, left=Inches(0.02), right=Inches(0.04),
                       top=Inches(0.02), bottom=Inches(0.02))

        # 구분선 (아주 옅은)
        divider_y = header_y + header_h + Inches(0.02)
        add_accent_bar(
            slide,
            cell.left + Inches(0.15), divider_y,
            cell.width - Inches(0.30), Inches(0.012),
            C_LGRAY,
        )

        # 라인 리스트 영역
        lines_y = divider_y + Inches(0.10)
        lines_w = cell.width - Inches(0.30)
        lines_x = cell.left + Inches(0.15)
        lines_h = cell.top + cell.height - lines_y - Inches(0.12)

        lines_tb = slide.shapes.add_textbox(
            lines_x, lines_y, lines_w, lines_h
        )
        lines_tf = lines_tb.text_frame
        lines_tf.word_wrap = True
        set_text_inset(lines_tb, left=Inches(0.02), right=Inches(0.02),
                       top=Inches(0.02), bottom=Inches(0.02))
        # 첫 단락 초기화
        lines_tf.paragraphs[0].text = ""

        for i, (label, value) in enumerate(card["lines"]):
            has_value = bool(value)
            # label이 "→" 나 "("로 시작하면 강조/부기 스타일
            is_arrow = label.startswith("→")
            is_paren = label.startswith("(") and label.endswith(")")

            if i == 0:
                p = lines_tf.paragraphs[0]
            else:
                p = lines_tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(4) if i > 0 else Pt(0)
            p.space_after = Pt(0)

            if has_value:
                # label : value
                r1 = p.add_run()
                r1.text = f"• {label}  "
                r1.font.size = Pt(10)
                r1.font.color.rgb = C_INK
                r1.font.bold = False

                r2 = p.add_run()
                r2.text = value
                r2.font.size = Pt(10)
                r2.font.color.rgb = card["accent"]
                r2.font.bold = True
            else:
                # 단일 텍스트 (화살표 결론 or 부기)
                r = p.add_run()
                r.text = label
                r.font.size = Pt(9.5 if is_paren else 10)
                if is_arrow:
                    r.font.color.rgb = card["accent"]
                    r.font.bold = True
                elif is_paren:
                    r.font.color.rgb = C_GRAY
                    r.font.italic = True
                else:
                    r.font.color.rgb = C_INK

    # ============================================================
    # 3) takeaway_bar — 카드 아래, 본문 최하단
    # ============================================================
    tb_y = body_bottom - takeaway_h
    tb_h = takeaway_h

    # 배경 패널 (navy 옅은 틴트)
    tb_panel = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, body_left, tb_y, body_width, tb_h
    )
    tb_panel.fill.solid()
    tb_panel.fill.fore_color.rgb = C_NAVY
    set_shape_opacity(tb_panel, 10)
    tb_panel.line.color.rgb = C_NAVY
    tb_panel.line.width = Pt(0.75)

    # 좌측 accent 바
    add_accent_bar(slide, body_left, tb_y, Inches(0.08), tb_h, C_NAVY)

    # 타이틀 + 3 bullet
    tb_inner_left = body_left + Inches(0.20)
    tb_inner_w = body_width - Inches(0.35)

    title_h = Inches(0.32)
    title_tb = add_textbox(
        slide, tb_inner_left, tb_y + Inches(0.06),
        tb_inner_w, title_h,
        "▶ 업계 주요 시사점",
        font_size=12, bold=True, color=C_NAVY,
        align=PP_ALIGN.LEFT,
    )
    set_text_inset(title_tb, left=Inches(0.04), right=Inches(0.04),
                   top=Inches(0.02), bottom=Inches(0.02))

    # 3 bullet을 단일 text_frame으로
    bullets_y = tb_y + Inches(0.06) + title_h
    bullets_h = tb_h - (bullets_y - tb_y) - Inches(0.06)

    bullets_tb = slide.shapes.add_textbox(
        tb_inner_left, bullets_y, tb_inner_w, bullets_h
    )
    bullets_tf = bullets_tb.text_frame
    bullets_tf.word_wrap = True
    set_text_inset(bullets_tb, left=Inches(0.04), right=Inches(0.04),
                   top=Inches(0.02), bottom=Inches(0.02))
    bullets_tf.paragraphs[0].text = ""

    takeaway_items = [
        [
            {"text": "도구: ", "font_size": 10.5, "color": C_INK, "bold": True},
            {"text": "GitHub Copilot · Azure OpenAI · Claude · 사내 LLM 프록시 (H-Chat) 4 가지 패턴",
             "font_size": 10.5, "color": C_INK},
        ],
        [
            {"text": "측정: ", "font_size": 10.5, "color": C_INK, "bold": True},
            {"text": "SPACE 5축 · 자체 설문 · Git 로그 분석 등 다각도 조합",
             "font_size": 10.5, "color": C_INK},
        ],
        [
            {"text": "2027.08 ", "font_size": 10.5, "color": C_RED, "bold": True},
            {"text": "EU AI Act 자동차 완전 적용 D-Day → 준비 시점",
             "font_size": 10.5, "color": C_INK, "bold": True},
        ],
    ]

    for i, segs in enumerate(takeaway_items):
        if i == 0:
            p = bullets_tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            # bullet dot
            r0 = p.add_run()
            r0.text = "• "
            r0.font.size = Pt(10.5)
            r0.font.color.rgb = C_NAVY
            r0.font.bold = True
            for seg in segs:
                r = p.add_run()
                r.text = seg["text"]
                r.font.size = Pt(seg["font_size"])
                r.font.color.rgb = seg["color"]
                r.font.bold = seg.get("bold", False)
        else:
            p = bullets_tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(2)
            p.space_after = Pt(0)
            r0 = p.add_run()
            r0.text = "• "
            r0.font.size = Pt(10.5)
            r0.font.color.rgb = C_NAVY
            r0.font.bold = True
            for seg in segs:
                r = p.add_run()
                r.text = seg["text"]
                r.font.size = Pt(seg["font_size"])
                r.font.color.rgb = seg["color"]
                r.font.bold = seg.get("bold", False)

    # ============================================================
    # 4) 각주 영역 (y=6.60~7.00) — 3줄, 9pt, gray
    # ============================================================
    footnote_y = Inches(6.60)
    footnote_h = Inches(0.40)

    entries = [
        "※ SPACE (Satisfaction·Performance·Activity·Communication·Efficiency): 개발자 생산성 측정 5축",
        "※ Faros AI (2025): AI 도구 다수 엔터프라이즈 분석 리포트",
        "※ EU AI Act: EU 인공지능 규제법. 2024.08 발효 / 2027.08 자동차 등 embedded high-risk AI 완전 적용",
    ]

    fn_tb = slide.shapes.add_textbox(
        body_left, footnote_y, body_width, footnote_h
    )
    fn_tf = fn_tb.text_frame
    fn_tf.word_wrap = True
    set_text_inset(fn_tb, left=Inches(0.02), right=Inches(0.02),
                   top=Inches(0.02), bottom=Inches(0.02))
    fn_tf.paragraphs[0].text = ""

    for i, entry in enumerate(entries):
        if i == 0:
            p = fn_tf.paragraphs[0]
        else:
            p = fn_tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.0
        run = p.add_run()
        run.text = entry
        run.font.size = Pt(9)
        run.font.color.rgb = C_GRAY
