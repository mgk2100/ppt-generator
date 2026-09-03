#!/usr/bin/env python3
"""SW 개발 AI (LLM) — Physical AI 포맷 6단계 파이프라인 1장 + 예상 Q&A 1장. 로보틱스 덱 마스터 사용.
사용: python3 build_deck.py [아이콘 폴더] [출력 pptx]   (기본 = mixed_icons → output/sw-dev-ai-llm.pptx)"""
import sys, os, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
BASE = f"{HERE}/base_master.pptx"          # 로보틱스 덱에서 슬라이드를 모두 제거한 마스터(144KB)
ICON_DIR = sys.argv[1] if len(sys.argv) > 1 else f"{HERE}/mixed_icons"
OUT = sys.argv[2] if len(sys.argv) > 2 else f"{REPO}/output/sw-dev-ai-llm.pptx"
SHOT = f"{REPO}/sources/sl-sw-agent/assets/ca_web_assist_conv.png"
FONT = "Noto Sans KR"

BLACK = RGBColor(0, 0, 0); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x20, 0x37, 0x68); BLUE = RGBColor(0x0E, 0x5D, 0xB7); BLUE2 = RGBColor(0x1F, 0x5F, 0xBF)
ARROW_BLUE = RGBColor(0x01, 0x34, 0x84); LABEL = RGBColor(0x00, 0x17, 0x61)
BORDER = RGBColor(0xA6, 0xCA, 0xEC); PANEL = RGBColor(0xF3, 0xF6, 0xFC); PANEL2 = RGBColor(0xE6, 0xEB, 0xF3)
CAPTION = RGBColor(0x20, 0x20, 0x20); GREY = RGBColor(0x59, 0x59, 0x59)
ARROW_FILL = RGBColor(0xB9, 0xD3, 0xEE); BENT = RGBColor(0x4C, 0x9B, 0xD6)

prs = Presentation(BASE)
layout = next(l for l in prs.slide_layouts if l.name == "1_제목 슬라이드")

# ---------- helpers ----------
def _font(run, size, bold=False, color=BLACK, italic=False):
    f = run.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = FONT
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)

def tb(slide, x, y, w, h, runs, size=9, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.MIDDLE, wrap=True, spacing=None, insets=(0, 0, 0, 0)):
    """runs: str 또는 [(text, {size,bold,color})...] 또는 [[...line1 runs...],[...line2...]]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(v) for v in insets]
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]
    for li, line in enumerate(runs):
        p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.line_spacing = spacing
        for text, st in line:
            r = p.add_run(); r.text = text
            _font(r, st.get("size", size), st.get("bold", bold), st.get("color", color))
    return box

def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=0.75, dash=None, adj=None):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(line_w)
        if dash: sh.line.dash_style = dash
    if adj is not None:
        for i, v in enumerate(adj): sh.adjustments[i] = v
    # 스타일 참조 제거(테마 효과 상속 방지)
    st = sh._element.find(qn("p:style"))
    if st is not None: sh._element.remove(st)
    return sh

def shape_text(sh, text, size, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, insets=(0, 0, 0, 0)):
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(v) for v in insets]
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; _font(r, size, bold, color)

def pic(slide, path, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)

def icon(slide, name, cx, cy, size=0.78):
    return pic(slide, f"{ICON_DIR}/{name}.png", cx - size / 2, cy - size / 2, w=size, h=size)

def small_arrow(slide, x, y):
    a = shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, 0.22, 0.21, fill=ARROW_FILL)
    return a

def clear_layout_placeholders(slide):
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)

def header(slide, tag, title, sub_runs):
    clear_layout_placeholders(slide)
    t = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.35, -0.21, 0.82, 0.82, fill=BLUE,
              line=RGBColor(0xC3, 0xC9, 0xD5), line_w=0.5, adj=[0.14561])
    shape_text(t, tag, 9, True, WHITE, insets=(0, 0.197, 0, 0))
    tb(slide, 1.26, 0.22, 9.05, 0.38, title, size=20, bold=True)
    tb(slide, 0.61, 0.70, 12.33, 0.50, sub_runs, size=16, bold=True)

def set_flipV(sh):
    sh._element.spPr.xfrm.set("flipV", "1")

# ---------- 슬라이드 1: 6단계 파이프라인 ----------
s1 = prs.slides.add_slide(layout)
header(s1, "LLM", "SW 개발 AI",
       [("설계자 질의·문서 자산화부터 사내 특화 LLM 까지, ", {}),
        ("6단계로 이어지는 LLM 개발 파이프라인", {"color": BLUE2})])

COLS = [0.36, 4.66, 8.97]; ROWS = [1.36, 3.44]; BW, BH = 4.01, 1.97
PX, PY, PW, PH = 0.14, 0.52, 3.73, 1.15   # 안쪽 패널 (박스 기준 오프셋)

STEPS = [
    dict(n="1", title="데이터 취득·자산화", key=True,
         items=[("설계자 ↔ AI 질의응답", "01_designer_chat"), ("문서·코드 업로드", "02_upload_docs"), ("데이터 자산화", "03_knowledge_db")],
         sub="질문·답변, 문서, 코드",
         caption="설계자의 질의·답변과 문서·코드를 지식 단위로 쌓는 과정"),
    dict(n="2", title="데이터 정제·구조화",
         items=[("정제·선별", "04_refine_funnel"), ("AI 학습 형태로 변환", "05_tokens"), ("학습 데이터", "16_dataset")],
         caption="중복·오류·민감정보를 걸러내 학습용 데이터로 바꾸는 단계"),
    dict(n="3", title="사내 LLM 학습",
         items=[("학습 데이터", "16_dataset"), ("범용 LLM", "06_neural_generic"), ("사내 적응 LLM", "07_neural_domain")],
         caption="사내 데이터를 범용 모델에 학습시켜 사내 LLM 을 만드는 단계"),
    dict(n="4", title="피드백 고도화",
         items=[("설계자 평가", "09_feedback"), ("사내 적응 LLM", "07_neural_domain"), ("고도화 LLM", "08_neural_advanced")],
         caption="설계자의 채택·수정 평가를 재학습해 정확도를 높이는 단계"),
    dict(n="5", title="경량화·배포",
         items=[("고도화 LLM", "08_neural_advanced"), ("지식증류 & 양자화", "10_distill_quantize"), ("사내 GPU 서버", "11_gpu_server")],
         caption="압축한 모델을 사내 서버에 올려 다수 설계자가 함께 쓰는 단계"),
    dict(n="6", title="활용·재수집",
         items=[("통합플랫폼 (웹·IDE)", "12_web_ide"), ("사내 LLM 답변", "17_llm_answer"), ("새 데이터 자산화", "03_knowledge_db")],
         sub="질문·답변, 문서, 코드",
         caption="플랫폼에서 답하며 생기는 새 데이터를 다시 자산으로 쌓는 단계"),
]

boxes = []
for i, st in enumerate(STEPS):
    bx = COLS[i % 3]; by = ROWS[i // 3]
    key = st.get("key", False)
    box = shape(s1, MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, BW, BH, fill=WHITE,
                line=(BLUE if key else BORDER), line_w=(1.5 if key else 0.75), adj=[0.035])
    boxes.append(box)
    num = shape(s1, MSO_SHAPE.OVAL, bx + 0.14, by + 0.14, 0.30, 0.30, fill=NAVY)
    shape_text(num, st["n"], 12, True, WHITE)
    tb(s1, bx + 0.54, by + 0.18, 3.33, 0.23, st["title"], size=13.5, bold=True, color=(BLUE if key else BLACK))
    panel = shape(s1, MSO_SHAPE.ROUNDED_RECTANGLE, bx + PX, by + PY, PW, PH, fill=PANEL, line=BORDER, line_w=1.0,
                  dash=MSO_LINE_DASH_STYLE.DASH, adj=[0.06])
    # 안쪽 3 항목
    px = bx + PX; py = by + PY
    slot_w = 1.10; gap = (PW - 3 * slot_w) / 2   # 슬롯 3개 + 화살표 자리
    for j, (label, ic) in enumerate(st["items"]):
        sx = px + j * (slot_w + gap); cx = sx + slot_w / 2
        tb(s1, sx - 0.08, py + 0.05, slot_w + 0.16, 0.21, label, size=8.5, bold=True, color=LABEL, align=PP_ALIGN.CENTER)
        has_sub = ("sub" in st) and j == 2
        icon(s1, ic, cx, py + 0.66 - (0.06 if has_sub else 0), size=0.66 if has_sub else 0.74)
        if has_sub:
            tb(s1, sx - 0.12, py + 0.95, slot_w + 0.24, 0.16, st["sub"], size=7, bold=True, color=LABEL, align=PP_ALIGN.CENTER)
        if j < 2:
            small_arrow(s1, sx + slot_w + gap / 2 - 0.11, py + 0.55)
    tb(s1, bx + 0.10, by + 1.70, 3.86, 0.26, st["caption"], size=9, color=CAPTION, anchor=MSO_ANCHOR.TOP)

# 박스 사이 → 화살표 (1→2, 2→3, 4→5, 5→6)
for r in ROWS:
    for x in (4.42, 8.73):
        tb(s1, x, r + 0.86, 0.20, 0.28, "→", size=14, bold=True, color=ARROW_BLUE, align=PP_ALIGN.CENTER)
# 3→4 는 되돌아가는 선으로 표현하지 않고, 우측 끝→좌측 시작 순환선(6→1) 만 표시
loop_pts = [(Inches(12.98), Inches(4.42)), (Inches(13.12), Inches(4.42)), (Inches(13.12), Inches(1.27)),
            (Inches(0.22), Inches(1.27)), (Inches(0.22), Inches(2.34)), (Inches(0.36), Inches(2.34))]
ff = s1.shapes.build_freeform(loop_pts[0][0], loop_pts[0][1])
ff.add_line_segments(loop_pts[1:], close=False)
loop = ff.convert_to_shape()
loop.fill.background(); loop.line.color.rgb = BLUE; loop.line.width = Pt(1.25)
ln = loop.line._get_or_add_ln(); tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
st_el = loop._element.find(qn("p:style"))
if st_el is not None: loop._element.remove(st_el)
tb(s1, 10.9, 1.08, 2.2, 0.17, "새 데이터 재수집 → 지속 학습", size=8, bold=True, color=BLUE, align=PP_ALIGN.RIGHT)

# 3→4 행 전환 화살표(오른쪽 위 → 왼쪽 아래)는 원본과 동일하게 생략(순환선이 대신 설명)

# ---------- 하단 좌: 실화면 + 굽은 화살표 ----------
shot = pic(s1, SHOT, 0.36, 5.55, w=1.80)
shot.line.color.rgb = BORDER; shot.line.width = Pt(0.75)
tb(s1, 0.36, 6.70, 1.80, 0.16, "SL SW Agent 웹 화면 (라이브)", size=7, color=GREY, align=PP_ALIGN.CENTER)
bent = shape(s1, MSO_SHAPE.BENT_ARROW, 2.23, 5.62, 4.01, 1.13, fill=BENT, line=RGBColor(0xC3, 0xC9, 0xD5), line_w=0.75,
             adj=[0.22666, 0.25, 0.25, 0.36479])
set_flipV(bent)
bent.adjustments[0] = 0.30   # 화살대 두께 확대 (2행 문구 수용)
tb(s1, 2.62, 6.265, 3.35, 0.40,
   [[("데이터 기반 지속 학습으로 상용 LLM 의존 축소,", {})], [("SL 특화 LLM 완성", {"bold": True})]],
   size=8.5, color=WHITE, align=PP_ALIGN.CENTER)

# ---------- 하단 우: 서비스 라이브러리 (Skill library 대응) ----------
LX, LY, LW, LH = 6.71, 5.52, 6.25, 1.90
lib = shape(s1, MSO_SHAPE.ROUNDED_RECTANGLE, LX, LY, LW, LH, fill=PANEL2, line=BORDER, line_w=0.75, adj=[0.035])
tb(s1, LX + 0.15, LY + 0.14, 3.6, 0.23, "SW 설계 지원 서비스", size=13.5, bold=True)
tb(s1, LX + LW - 2.4, LY + 0.16, 2.25, 0.21, "통합플랫폼 (SL SW Agent)", size=9, bold=True, color=LABEL, align=PP_ALIGN.RIGHT)
inner = shape(s1, MSO_SHAPE.ROUNDED_RECTANGLE, LX + 0.14, LY + 0.47, LW - 0.28, 1.06, fill=PANEL, line=BORDER, line_w=1.0,
              dash=MSO_LINE_DASH_STYLE.DASH, adj=[0.06])
SERVICES = [
    ("문서 생성", "13_doc_gen", "자산: 소스코드·설계문서", "활용: SAD·SDD 자동 작성"),
    ("코드 생성", "14_code_gen", "자산: 질의·답변·사내 코드", "활용: 코드 작성·질의응답"),
    ("정적 검증", "15_static_check", "자산: 위배 사례·수정안", "활용: 코딩 규칙 검사·수정안"),
]
cw = (LW - 0.28 - 0.2) / 3
for k, (name, ic, a1, a2) in enumerate(SERVICES):
    cx0 = LX + 0.14 + 0.1 + k * cw
    card = shape(s1, MSO_SHAPE.ROUNDED_RECTANGLE, cx0 + 0.04, LY + 0.55, cw - 0.08, 0.90, fill=WHITE, line=BORDER, line_w=0.75, adj=[0.06])
    icon(s1, ic, cx0 + 0.36, LY + 0.99, size=0.50)
    tb(s1, cx0 + 0.66, LY + 0.60, cw - 0.72, 0.22, name, size=10, bold=True, color=LABEL)
    tb(s1, cx0 + 0.66, LY + 0.83, cw - 0.70, 0.56,
       [[(a1, {})], [(a2, {})]], size=7.2, color=CAPTION, anchor=MSO_ANCHOR.TOP)
    if k < 2:
        small_arrow(s1, cx0 + cw - 0.13, LY + 0.90)
tb(s1, LX + 0.14, LY + 1.58, LW - 0.28, 0.19,
   "세 서비스가 한 플랫폼에서 지식 자산과 LLM 을 공유. 정적 검증 → 문서 생성 → 코드 생성 순으로 사내 LLM 비중 확대",
   size=8.6, color=CAPTION)

# 각주 (좌하단)
tb(s1, 0.36, 6.98, 6.3, 0.40,
   [[("※ LLM(Large Language Model): 대규모 언어 모델 · 지식증류: 고성능 AI 답변으로 소형 AI 를 학습시키는 기법 · 양자화: 모델 압축 기법", {})],
    [("※ SAD·SDD: SW 아키텍처·상세 설계 문서 · MISRA(Motor Industry Software Reliability Association): 차량용 C 코딩 규칙", {})]],
   size=6.8, color=GREY, anchor=MSO_ANCHOR.TOP)

# ---------- 슬라이드 2: 예상 Q&A ----------
s2 = prs.slides.add_slide(layout)
header(s2, "Q&A", "SW 개발 AI — 예상 질의 대응",
       [("임원 보고 시 나올 수 있는 질문과 핵심 답변·근거 ", {}), ("(팀장 대응용)", {"color": BLUE2})])

QA = [
    ("데이터는 지금 얼마나 확보돼 있나?",
     "사내 코드 약 25만 조각(약 2천만 토큰) 확보. 설계자 질의·답변은 약 100건으로 아직 미미 → 플랫폼 오픈(2027.1) 후 본격 축적.",
     "2026.07 실측. 데이터 취득이 현재 가장 부족한 역량"),
    ("데이터 취득은 어떻게 늘리나?",
     "설계자가 상용 LLM 을 쓰는 순간 질문·답변·평가(채택/수정)가 자동 기록되고, 업로드 문서·코드는 자동 자산화. 대화를 지식 단위(질답·결정·패턴·주의점 등 7종)로 추출하는 기능 구현 완료.",
     "기록·피드백 기능 구축 2026.07, 지식 추출 파이프라인 구현 2026.07"),
    ("서버 스펙과 비용은?",
     "Dell R770 2대: GPU RTX PRO 6000 96GB 총 3장(VRAM 288GB), 메모리 256GB·512GB, SSD 3.84TB×2·7.68TB×2. 2대 합계 약 8,000만원. 1대 서비스(추론), 1대 학습·실험.",
     "(주)아이웍스 도입, 2026.08 확정 스펙"),
    ("학습에 얼마나 걸리나?",
     "소형 모델 실측: 3.5만 건 4회 반복 학습에 약 2시간(GPU 1장, 메모리 52GB). 30B급 사내 특화 학습(2천만 토큰·3회 반복)은 GPU 1장 기준 하루 안팎으로 추정. 병목은 학습 시간이 아니라 학습 데이터 확보.",
     "실측 2026.07 학습 기록 / 30B급은 추정치(완료된 학습 이력 없음)"),
    ("사내 LLM 이 상용 수준이 되나?",
     "15종 이상 실서버 실측 결과 96GB GPU 로는 30B급이 현실적 상한이며 상용 프론티어 대비 성능 격차 존재. 초기는 상용 중심, 데이터 축적 후 지식증류로 격차를 좁혀 쉬운 영역(정적 검증→문서 생성)부터 병행.",
     "실측 2026.06~07. 현재 사내 모델은 코드 검색·분석 서술 등 한정 용도"),
    ("상용 LLM 비용은?",
     "코딩 지원·정적 검증 리뷰 인당 월 약 1.5만원, 프로젝트 심층 분석 건당 약 15만원(38만 줄 기준), 문서 생성은 구독형 경로 사용. 사내 LLM 병행 시 단계적으로 절감.",
     "2026.03 계획서 단가(Enterprise API 기준)"),
    ("최상위 모델을 자체 서버에 올리면?",
     "최상위 오픈 모델(Kimi K3) 기준 GPU 약 20장 ≈ 3.7억원, 서버 4~5대. 보유 GPU 는 카드 간 고속 연결(NVLink) 미지원이라 데이터센터급 장비가 별도 필요 → 현 예산에선 비현실적.",
     "2026.07 산정(RTX PRO 6000 1,850만원/장)"),
    ("사내 코드가 외부로 나가지 않나?",
     "사내 코드를 외부 학습 서비스로 보내는 경로는 기본 차단(보안 게이트). 상용 LLM 은 기업용 계약 경로로만 사용, 사내 LLM 학습·추론은 사내 서버 안에서 처리. 상용 답변의 학습 활용은 약관 확인 병행.",
     "보안 게이트 구현 2026.06, 법무 확인 권고 2026.07"),
    ("일정은?",
     "2026.12 테스터 기반 1차 통합 구축 → 2027.1 중순 SW 설계자 오픈 → 2027.1 부터 축적 데이터로 사내 LLM 1차 학습 착수.",
     "2026.07 확정 로드맵"),
    ("기대 효과는?",
     "설계 문서(SAD·SDD) 작성 공수 80% 이상 절감 목표(컴포넌트 1개 문서 약 50분 자동 생성), 2,040개 파일 정적 검증 35분 완주, 신규 인력 숙련 시간 단축.",
     "목표치 2026.03 계획서 / 소요 시간은 실측"),
]
TX, TY, TW = 0.36, 1.36, 12.62
colw = [2.55, 6.75, 3.32]
rows_n = len(QA) + 1
tbl_shape = s2.shapes.add_table(rows_n, 3, Inches(TX), Inches(TY), Inches(TW), Inches(0.42 * rows_n))
tbl = tbl_shape.table
tblPr = tbl._tbl.tblPr
tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
sid = tblPr.find(qn("a:tableStyleId"))
if sid is not None: tblPr.remove(sid)
for c, w in enumerate(colw): tbl.columns[c].width = Inches(w)

def cell_set(cell, text, size, bold=False, color=BLACK, fill=None, align=PP_ALIGN.LEFT):
    cell.text = ""
    tf = cell.text_frame; tf.word_wrap = True
    cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.08); cell.margin_top = Inches(0.05); cell.margin_bottom = Inches(0.05)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; _font(r, size, bold, color)
    tcPr = cell._tc.get_or_add_tcPr()
    for el in tcPr.findall(qn("a:solidFill")) + tcPr.findall(qn("a:noFill")): tcPr.remove(el)
    if fill is not None:
        sf = etree.SubElement(tcPr, qn("a:solidFill")); clr = etree.SubElement(sf, qn("a:srgbClr")); clr.set("val", str(fill))
    else:
        etree.SubElement(tcPr, qn("a:noFill"))
    # 테두리: 아래쪽 얇은 선
    for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(edge)): tcPr.remove(el)
    for edge, col, w in (("a:lnL", None, 0), ("a:lnR", None, 0), ("a:lnT", "A6CAEC", 6350), ("a:lnB", "A6CAEC", 6350)):
        ln_el = etree.SubElement(tcPr, qn(edge))
        if col is None:
            ln_el.set("w", "0"); etree.SubElement(ln_el, qn("a:noFill"))
        else:
            ln_el.set("w", str(w)); sf = etree.SubElement(ln_el, qn("a:solidFill")); c2 = etree.SubElement(sf, qn("a:srgbClr")); c2.set("val", col)
    # OOXML 순서: lnL lnR lnT lnB ... fill  → fill 요소를 맨 뒤로 이동
    for tag in ("a:solidFill", "a:noFill"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el); tcPr.append(el)

for c, h in enumerate(["예상 질문", "핵심 답변", "근거 · 비고"]):
    cell_set(tbl.cell(0, c), h, 10, True, WHITE, fill="203768", align=PP_ALIGN.CENTER)
tbl.rows[0].height = Inches(0.34)
for r, (q, a, b) in enumerate(QA, start=1):
    zebra = "F3F6FC" if r % 2 == 0 else None
    cell_set(tbl.cell(r, 0), f"Q{r}. {q}", 9, True, NAVY, fill=zebra)
    cell_set(tbl.cell(r, 1), a, 8.6, False, BLACK, fill=zebra)
    cell_set(tbl.cell(r, 2), b, 8, False, GREY, fill=zebra)
    tbl.rows[r].height = Inches(0.50)

tb(s2, 0.36, 7.02, 12.6, 0.40,
   [[("※ '추정' 표기가 없는 수치는 실측·확정값.  지식증류: 고성능 AI 의 답변으로 소형 AI 를 학습시키는 기법  ·  토큰: AI 가 글을 처리하는 최소 단위(한글 약 1~2자)  ·  30B: 모델 크기(매개변수 300억 개)", {})],
    [("※ SAD(Software Architecture Design Specification)·SDD(Software Detailed Design Specification): SW 아키텍처·상세 설계 문서  ·  NVLink: GPU 간 고속 직결 연결  ·  MISRA(Motor Industry Software Reliability Association): 차량용 C 코딩 규칙", {})]],
   size=7, color=GREY, anchor=MSO_ANCHOR.TOP)

# ---------- 저장 ----------
prs.save(OUT)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")
