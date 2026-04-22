"""Slide 5 — 실행 로드맵 및 승인 요청 사항.

구성:
1) 상단: Key Message Bar
2) 중단: CHEVRON 3단계 로드맵 (단기/중기/장기)
3) 하단 2열 분할: 좌) 자사 시범 사업 프로토콜 · 우) 승인 요청 사항

제약:
- 자사 자원·기간 구체 수치 없음 ("별도 산정" 표기)
- Phase 기간은 "단기/중기/장기"만 사용
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from ppt_utils import (
    set_title,
    clear_placeholders,
    add_textbox,
    add_para,
    add_rich_text,
    add_accent_bar,
    make_icon_circle,
    make_icon_badge,
    add_shadow,
    set_body_anchor,
    set_text_inset,
    calc_grid,
)
from template_contract import CONTENT_SAFE, LAYOUT_CONTENT


# ============================ 색상 팔레트 ============================
C_NAVY   = RGBColor(0x1F, 0x49, 0x7D)
C_BLUE   = RGBColor(0x4F, 0x81, 0xBD)
C_GREEN  = RGBColor(0x2E, 0x8B, 0x57)
C_TEAL   = RGBColor(0x2C, 0x7F, 0x94)
C_INK    = RGBColor(0x1A, 0x1F, 0x2E)
C_GRAY   = RGBColor(0x66, 0x66, 0x66)
C_LGRAY  = RGBColor(0xE0, 0xE4, 0xEA)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_OFFWHITE = RGBColor(0xF7, 0xF9, 0xFC)
BG_BLUE  = RGBColor(0xE7, 0xF0, 0xFA)
BG_TEAL  = RGBColor(0xE3, 0xF2, 0xF7)


# ============================ 헬퍼 ============================

def _add_chevron_step(slide, x, y, w, h, phase, body, fill_color):
    """CHEVRON 한 단계: 상단 phase(볼드) + 하단 body(2줄)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    add_shadow(shape, blur_pt=4, dist_pt=2, opacity_pct=28, color=C_GRAY)

    tf = shape.text_frame
    tf.word_wrap = True
    set_text_inset(shape, left=Inches(0.14), top=Inches(0.08),
                   right=Inches(0.26), bottom=Inches(0.08))
    set_body_anchor(shape, "ctr")

    # 1단락: phase 라벨 (볼드)
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    run0 = p0.add_run()
    run0.text = phase
    run0.font.size = Pt(11.5)
    run0.font.bold = True
    run0.font.color.rgb = C_WHITE

    # 2~3단락: body (줄바꿈 분리)
    body_lines = body.split("\n")
    for i, line in enumerate(body_lines):
        add_para(
            tf, line,
            font_size=9.5, color=C_WHITE, bold=False,
            align=PP_ALIGN.CENTER,
            space_before=Pt(2) if i == 0 else Pt(1),
        )
    return shape


def _add_info_card(slide, x, y, w, h, title, items, accent_color, bg_color):
    """정보 카드: BG 배경 + 상단 accent 바 + 제목 + key:value 리스트."""
    # 본체
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    body.adjustments[0] = 0.04
    body.fill.solid()
    body.fill.fore_color.rgb = bg_color
    body.line.color.rgb = C_LGRAY
    body.line.width = Pt(0.75)
    add_shadow(body, blur_pt=4, dist_pt=2, opacity_pct=18, color=C_GRAY)

    # 상단 accent 바
    bar_h = Inches(0.08)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, x + Inches(0.02), y + Inches(0.02),
        w - Inches(0.04), bar_h,
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_color
    bar.line.fill.background()

    # 카드 제목
    pad_l = Inches(0.18)
    pad_r = Inches(0.18)
    title_x = x + pad_l
    title_y = y + Inches(0.16)
    title_w = w - pad_l - pad_r
    title_h = Inches(0.36)
    title_box = add_textbox(
        slide, title_x, title_y, title_w, title_h, title,
        font_size=12.5, color=accent_color, bold=True, align=PP_ALIGN.LEFT,
    )
    set_text_inset(title_box, left=Inches(0.02), top=Inches(0.02),
                   right=Inches(0.02), bottom=Inches(0.02))

    # key:value 리스트 영역
    list_x = x + pad_l
    list_y = title_y + title_h + Inches(0.06)
    list_w = w - pad_l - pad_r
    list_bottom = y + h - Inches(0.14)
    list_h = list_bottom - list_y

    list_box = slide.shapes.add_textbox(list_x, list_y, list_w, list_h)
    tf = list_box.text_frame
    tf.word_wrap = True
    set_text_inset(list_box, left=Inches(0.02), top=Inches(0.02),
                   right=Inches(0.02), bottom=Inches(0.02))

    for i, item in enumerate(items):
        key = item["key"]
        value = item["value"]
        if i == 0:
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
        else:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(3)

        # key (볼드, accent)
        rk = p.add_run()
        rk.text = f"{key}  "
        rk.font.size = Pt(9.5)
        rk.font.bold = True
        rk.font.color.rgb = accent_color

        # 구분
        rsep = p.add_run()
        rsep.text = "·  "
        rsep.font.size = Pt(9.5)
        rsep.font.color.rgb = C_GRAY

        # value
        rv = p.add_run()
        rv.text = value
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = C_INK

    return body


def _add_approval_card(slide, x, y, w, h, title, items, accent_color, bg_color):
    """승인 요청 카드: 번호 배지 + 라벨(볼드) + note(회색)."""
    # 본체
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    body.adjustments[0] = 0.04
    body.fill.solid()
    body.fill.fore_color.rgb = bg_color
    body.line.color.rgb = C_LGRAY
    body.line.width = Pt(0.75)
    add_shadow(body, blur_pt=4, dist_pt=2, opacity_pct=18, color=C_GRAY)

    # 상단 accent 바
    bar_h = Inches(0.08)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, x + Inches(0.02), y + Inches(0.02),
        w - Inches(0.04), bar_h,
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_color
    bar.line.fill.background()

    # 카드 제목
    pad_l = Inches(0.18)
    pad_r = Inches(0.18)
    title_x = x + pad_l
    title_y = y + Inches(0.16)
    title_w = w - pad_l - pad_r
    title_h = Inches(0.36)
    title_box = add_textbox(
        slide, title_x, title_y, title_w, title_h, title,
        font_size=12.5, color=accent_color, bold=True, align=PP_ALIGN.LEFT,
    )
    set_text_inset(title_box, left=Inches(0.02), top=Inches(0.02),
                   right=Inches(0.02), bottom=Inches(0.02))

    # 아이템 리스트 — 번호 배지 원형 + 텍스트 영역
    list_top = title_y + title_h + Inches(0.08)
    list_bottom = y + h - Inches(0.14)
    list_h = list_bottom - list_top
    n_items = len(items)
    row_h = list_h // n_items if n_items > 0 else list_h

    badge_size = Inches(0.28)
    for i, item in enumerate(items):
        row_y = list_top + i * row_h

        # 번호 배지 (원형)
        badge_x = x + pad_l
        badge_y = row_y + (row_h - badge_size) // 2
        make_icon_circle(
            slide, badge_x, badge_y, badge_size,
            accent_color, text=item["num"],
            font_size=9, font_color=C_WHITE,
        )

        # 텍스트 영역 (라벨 + note)
        text_x = badge_x + badge_size + Inches(0.08)
        text_w = (x + w - pad_r) - text_x
        text_box = slide.shapes.add_textbox(text_x, row_y, text_w, row_h)
        tf = text_box.text_frame
        tf.word_wrap = True
        set_text_inset(text_box, left=Inches(0.02), top=Inches(0.02),
                       right=Inches(0.02), bottom=Inches(0.02))
        set_body_anchor(text_box, "ctr")

        # 라벨 (볼드)
        p0 = tf.paragraphs[0]
        p0.alignment = PP_ALIGN.LEFT
        rl = p0.add_run()
        rl.text = item["label"]
        rl.font.size = Pt(10)
        rl.font.bold = True
        rl.font.color.rgb = C_INK

        # note (회색, 작게)
        add_para(
            tf, item["note"],
            font_size=8.5, color=C_GRAY, bold=False,
            align=PP_ALIGN.LEFT, space_before=Pt(1),
        )

    return body


# ============================ 메인 빌더 ============================

def build_slide_5(slide):
    set_title(slide, "실행 로드맵 및 승인 요청 사항")
    clear_placeholders(slide, keep=[0])

    # CONTENT_SAFE 좌표
    safe_left = CONTENT_SAFE.left
    safe_top = CONTENT_SAFE.top
    safe_w = CONTENT_SAFE.width
    safe_bottom = CONTENT_SAFE.bottom

    # ------------------------------------------------------------------
    # 1) Key Message Bar  (상단)
    # ------------------------------------------------------------------
    km_x = safe_left
    km_y = safe_top + Inches(0.02)           # ~0.70"
    km_w = safe_w
    km_h = Inches(0.48)

    # 좌측 accent 세로 바
    accent_bar_w = Inches(0.08)
    add_accent_bar(slide, km_x, km_y, accent_bar_w, km_h, C_TEAL)

    # 본체 박스 (연한 배경)
    km_body_x = km_x + accent_bar_w
    km_body_w = km_w - accent_bar_w
    km_body = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, km_body_x, km_y, km_body_w, km_h,
    )
    km_body.fill.solid()
    km_body.fill.fore_color.rgb = C_OFFWHITE
    km_body.line.color.rgb = C_LGRAY
    km_body.line.width = Pt(0.5)

    # KEY MESSAGE 배지
    km_badge_w = Inches(1.10)
    km_badge_h = Inches(0.32)
    km_badge_x = km_body_x + Inches(0.10)
    km_badge_y = km_y + (km_h - km_badge_h) // 2
    make_icon_badge(
        slide, km_badge_x, km_badge_y, km_badge_w, km_badge_h,
        "KEY MESSAGE", C_TEAL, font_size=9, font_color=C_WHITE,
        corner_radius=0.25,
    )

    # 메시지 텍스트
    km_text_x = km_badge_x + km_badge_w + Inches(0.15)
    km_text_w = km_body_x + km_body_w - km_text_x - Inches(0.10)
    km_text = slide.shapes.add_textbox(
        km_text_x, km_y, km_text_w, km_h,
    )
    km_tf = km_text.text_frame
    km_tf.word_wrap = True
    set_text_inset(km_text, left=Inches(0.02), top=Inches(0.04),
                   right=Inches(0.02), bottom=Inches(0.04))
    set_body_anchor(km_text, "ctr")
    p_km = km_tf.paragraphs[0]
    p_km.alignment = PP_ALIGN.LEFT
    r1 = p_km.add_run()
    r1.text = "단계적 도입"
    r1.font.size = Pt(11); r1.font.bold = True; r1.font.color.rgb = C_GREEN
    r2 = p_km.add_run()
    r2.text = " + "
    r2.font.size = Pt(11); r2.font.color.rgb = C_GRAY
    r3 = p_km.add_run()
    r3.text = "자사 시범 사업"
    r3.font.size = Pt(11); r3.font.bold = True; r3.font.color.rgb = C_BLUE
    r4 = p_km.add_run()
    r4.text = "으로 실측  ·  "
    r4.font.size = Pt(11); r4.font.color.rgb = C_INK
    r5 = p_km.add_run()
    r5.text = "선행 영역 우선 착수"
    r5.font.size = Pt(11); r5.font.bold = True; r5.font.color.rgb = C_NAVY

    # ------------------------------------------------------------------
    # 2) CHEVRON 3단계 로드맵 (중단)
    # ------------------------------------------------------------------
    rd_section_y = km_y + km_h + Inches(0.14)   # ~1.32"

    # 섹션 소제목
    rd_title_h = Inches(0.28)
    rd_title = add_textbox(
        slide, safe_left, rd_section_y, Inches(5.5), rd_title_h,
        "① 단계적 도입 로드맵 (업계 일반 권고)",
        font_size=11, color=C_NAVY, bold=True,
    )

    # 우측 보조 캡션
    cap_w = Inches(4.2)
    cap_x = safe_left + safe_w - cap_w
    cap = add_textbox(
        slide, cap_x, rd_section_y, cap_w, rd_title_h,
        "기간 구체화는 자사 시범 사업 후 확정",
        font_size=9, color=C_GRAY, bold=False, align=PP_ALIGN.RIGHT,
    )

    # CHEVRON 배치 영역
    ch_y = rd_section_y + rd_title_h + Inches(0.06)  # ~1.66"
    ch_h = Inches(0.90)
    ch_area_x = safe_left
    ch_area_w = safe_w
    overlap = Inches(0.14)

    phases = [
        ("Phase 1  ·  단기",
         "도구·테스트 자동화 파일럿\n위험 낮음 · 효과 큼",
         C_GREEN),
        ("Phase 2  ·  중기",
         "선행 R&D · 응용SW 확대\n사내 AI 학습 · 자동 검사 통합",
         C_BLUE),
        ("Phase 3  ·  장기",
         "양산 비안전 영역 단계 적용\n양산 안전은 시범 측정 후",
         C_NAVY),
    ]
    n = len(phases)
    step_w = (ch_area_w + (n - 1) * overlap) // n

    for i, (phase, body, color) in enumerate(phases):
        cx = ch_area_x + i * (step_w - overlap)
        _add_chevron_step(slide, cx, ch_y, step_w, ch_h, phase, body, color)

    # ------------------------------------------------------------------
    # 3) 하단 2열 (좌: 시범 사업 · 우: 승인 요청)
    # ------------------------------------------------------------------
    bt_top = ch_y + ch_h + Inches(0.22)          # ~2.78"
    bt_bottom = safe_bottom - Inches(0.04)       # ~6.98"
    bt_h = bt_bottom - bt_top                    # ~4.20"
    bt_area = (safe_left, bt_top, safe_w, bt_h)

    bt_grid = calc_grid(
        rows=1, cols=2, area=bt_area,
        gap=Inches(0.22),
    )

    # 좌: 자사 시범 사업 (실측 프로토콜)
    left_cell = bt_grid[0][0]
    left_items = [
        {"key": "대상",       "value": "영역별 대표 SW 모듈 (선행 / 양산 일반 / 양산 안전 / 검증 자동화)"},
        {"key": "기간",       "value": "별도 산정"},
        {"key": "규모·예산",  "value": "별도 산정 (파일럿 계획 수립 단계)"},
        {"key": "측정 지표",  "value": "공수 · 결함밀도 · 코드 리뷰 부담"},
        {"key": "목적",       "value": "자사 환경에서 영역별 실효성 실측"},
    ]
    _add_info_card(
        slide,
        left_cell.left, left_cell.top, left_cell.width, left_cell.height,
        "②  자사 시범 사업 (실측 프로토콜)",
        left_items, C_BLUE, BG_BLUE,
    )

    # 우: 승인 요청 사항
    right_cell = bt_grid[0][1]
    right_items = [
        {"num": "①", "label": "선행 · 테스트 자동화 영역 도입 검토 착수",
         "note": "효과 입증된 영역"},
        {"num": "②", "label": "자사 시범 사업 추진 승인",
         "note": "규모·예산 별도 산정"},
        {"num": "③", "label": "품질검증 인력 보강 계획 수립",
         "note": "코드 리뷰 부담 증가 대응"},
        {"num": "④", "label": "사내 AI 사용 가이드라인 제정 착수",
         "note": "EU AI Act 2027.08 대응"},
    ]
    _add_approval_card(
        slide,
        right_cell.left, right_cell.top, right_cell.width, right_cell.height,
        "③  승인 요청 사항",
        right_items, C_TEAL, BG_TEAL,
    )
