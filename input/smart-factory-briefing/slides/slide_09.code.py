"""slide_09 — 운영 경험 및 도입 시 참고 사항 (pattern: table).

디자인 시스템: input/smart-factory-briefing/design_system.md (Electric Blue + Ink,
맑은 고딕 단일(force_font 일괄 적용), 그림자·이모지·zebra 금지, 수평 hairline 위주).
톤: 참고 조 — 상대 계획 수치 인용 없음, 저희 실측 경험만 서술.
"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from ppt_utils import (
    set_title, clear_placeholders,
    add_grid_table, set_cell_border,
    set_body_anchor, set_text_inset,
    estimate_container_height,
    force_font,
)
from template_contract import CONTENT_SAFE


# ---------------- 디자인 토큰 (design_system.md — 이 목록 밖 색 금지) ----------------
PRIMARY        = RGBColor(0x02, 0x4A, 0xD8)  # Electric Blue — 유일한 신호색
PRIMARY_BRIGHT = RGBColor(0x29, 0x6E, 0xF9)  # 다크 슬랩 위 강조 전용
INK            = RGBColor(0x1A, 0x1A, 0x1A)
CHARCOAL       = RGBColor(0x3D, 0x3D, 0x3D)
CANVAS         = RGBColor(0xFF, 0xFF, 0xFF)
CLOUD          = RGBColor(0xF7, 0xF7, 0xF7)
FOG            = RGBColor(0xE8, 0xE8, 0xE8)
STEEL          = RGBColor(0xC2, 0xC2, 0xC2)
ON_INK         = RGBColor(0xFF, 0xFF, 0xFF)

FONT_R  = "Pretendard"
FONT_SB = "Pretendard SemiBold"

_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ---------------- 내부 헬퍼 ----------------

def _set_ea_font(run, name):
    """run의 East Asian typeface를 latin과 동일하게 — 한글도 Pretendard로 렌더."""
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(_A_NS + "ea")
    if ea is None:
        ea = rPr.makeelement(_A_NS + "ea", {})
        latin = rPr.find(_A_NS + "latin")
        if latin is not None:
            latin.addnext(ea)  # rPr 자식 순서: latin 바로 뒤가 ea 자리
        else:
            rPr.append(ea)
    ea.set("typeface", name)


def _ea_pass_tf(tf):
    for p in tf.paragraphs:
        for run in p.runs:
            if run.font.name:
                _set_ea_font(run, run.font.name)


def _finalize_korean_typography(slide):
    """슬라이드 전체 run에 EA 폰트 반영 (제목 플레이스홀더·표 셀 포함)."""
    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    _ea_pass_tf(cell.text_frame)
        elif shape.has_text_frame:
            _ea_pass_tf(shape.text_frame)


def _plain(content):
    """segments 리스트 → 평문 (행 높이 추정용)."""
    if isinstance(content, str):
        return content
    return "".join(s if isinstance(s, str) else s["text"] for s in content)


def _sb(text):
    """'도입 시 참고' 열 핵심 구절 — SemiBold INK 세그먼트."""
    return {"text": text, "color": INK, "font_name": FONT_SB,
            "font_size": 9.5, "bold": False}


# ---------------- 콘텐츠 (spec content_blocks — 임의 추가 없음) ----------------

LEAD_SEGS = [
    ("서버 구축 · 운영에서 확인한 ", INK),
    ("5가지", PRIMARY),
    (" — 도입 판단에 참고할 사항", INK),
]

HEADERS = ["구분", "실측", "도입 시 참고"]

# (구분, 실측, 도입 시 참고[segments])
ROWS = [
    ("저장 용량",
     "모델 저장소만 1.4TB 사용 중\n"
     "30B급 20~60GB · 100B급+ 130~180GB, 비교용 모델 누적 빠름",
     ["모델 1~2TB + 학습 데이터를 빼면 제조 데이터까지 같은 서버에 두기엔 "
      "남는 공간 부족 — ",
      _sb("대용량 소수 구성"),
      "(예: 7.68TB × 2~4)이 슬롯 · 관리 유리"]),
    ("서버 역할 분리",
     "AI 환경 업데이트 재부팅마다 DB 동반 다운 · 학습 자원 점유 시 "
     "DB 응답 저하 → 분리 운영",
     ["GPU 서버와 데이터 플랫폼 서버 ",
      _sb("분리 구성"),
      " (데이터 서버에 GPU 불요)"]),
    ("전원 · 발열",
     "RTX PRO 6000 제품군별 소비전력 300W / 600W 2배 차이 — "
     "600W × 4장 = GPU만 2,400W",
     ["일반 사무실 콘센트 불가 — ",
      _sb("전용 회로 · 전산실 설치"),
      ". 견적 시 제품군 확인"]),
    ("운영체제",
     "Ubuntu 단일 운영 — AI 도구 대부분 Linux 기준",
     ["듀얼부팅은 OS 전환 시 전원 차단 필요 → ",
      _sb("공용 서버 부적합")]),
    ("GPU 사용률 기록",
     "초기 기록 부재로 배치 수 차례 조정",
     ["도입 초기부터 ",
      _sb("사용률 기록"),
      " → 증설 판단 근거 확보"]),
]

CLOSING_TEXT_HEAD = "저희 구축 과정에서 확인한 내용 공유 — 필요 시 도입처 담당자 연결 등 "
CLOSING_TEXT_EMPH = "지원 가능"


# ---------------- 빌드 ----------------

def build_slide_9(slide):
    set_title(slide, "운영 경험 및 도입 시 참고 사항", font_size=20, bold=True)
    clear_placeholders(slide, keep=[0])

    x0 = CONTENT_SAFE.left
    full_w = CONTENT_SAFE.width

    # ---- 1. 리드 (표준 규격 2-1: y=0.72" h=0.35" 15pt bold INK — '5가지' PRIMARY) ----
    lead_top = Inches(0.72)
    lead_h = Inches(0.35)
    lead_tb = slide.shapes.add_textbox(x0, lead_top, full_w, lead_h)
    tf = lead_tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for text, color in LEAD_SEGS:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = color

    # ---- 3. 하단 다크 슬랩 좌표 먼저 확정 (표 가용 높이 산출용) ----
    slab_h = Inches(0.5)
    slab_top = CONTENT_SAFE.top + CONTENT_SAFE.height - slab_h  # bottom = 7.02"

    # ---- 2. 메인 표 add_grid_table — 3열 [구분 | 실측 | 도입 시 참고] ----
    table_top = max(lead_top + lead_h + Inches(0.12), Inches(1.18))  # 본문 시작 y ≥ 1.18"
    table_bottom_max = slab_top - Inches(0.15)
    avail_h = table_bottom_max - table_top

    col_ratios = [1.3, 4.4, 4.4]
    ratio_sum = sum(col_ratios)
    col_ws = [int(full_w * r / ratio_sum) for r in col_ratios]
    # 셀 내부 좌우 마진(기본 0.1"×2) + 여유 → 텍스트 가용 폭
    text_ws = [cw - Inches(0.24) for cw in col_ws]

    # 행 높이 = 3개 셀 중 최대 필요 높이 (estimate_container_height 기반 — 넘침 금지)
    # 맑은 고딕은 미리보기 폰트보다 실폭이 넓어 줄 수가 늘 수 있음 → 높이 여유 +15% 가산
    ROW_H_GAIN = 1.15
    row_hs = []
    for gubun, exp, ref in ROWS:
        est = max(
            estimate_container_height(gubun, 9.5, max_width=text_ws[0],
                                      padding_top=Inches(0.10),
                                      padding_bottom=Inches(0.10)),
            estimate_container_height(exp, 9.5, max_width=text_ws[1],
                                      padding_top=Inches(0.10),
                                      padding_bottom=Inches(0.10)),
            estimate_container_height(_plain(ref), 9.5, max_width=text_ws[2],
                                      padding_top=Inches(0.10),
                                      padding_bottom=Inches(0.10)),
        )
        row_hs.append(max(int(est * ROW_H_GAIN), int(Inches(0.5))))

    header_h = int(Inches(0.40))
    total_h = header_h + sum(row_hs)
    if total_h < avail_h:
        # 남는 공간을 행에 균등 배분 → 하단 빈 영역 없이 표가 영역을 채움
        extra = int(avail_h - total_h) // len(row_hs)
        row_hs = [h + extra for h in row_hs]
    elif total_h > avail_h:
        # 방어적 클램프 (estimator 여유분 내 비례 축소)
        scale = float(avail_h - header_h) / float(total_h - header_h)
        row_hs = [int(h * scale) for h in row_hs]
    table_h = header_h + sum(row_hs)

    body_border = dict(border_edges="b", border_color=FOG, border_width_pt=0.75)
    cells = {}
    for c, htxt in enumerate(HEADERS):
        cells[(0, c)] = dict(
            text=htxt, fill=CLOUD, font_name=FONT_SB, font_size=10.5,
            font_color=INK, bold=False, align="l", anchor="ctr",
            border_edges="b", border_color=STEEL, border_width_pt=1.0,
        )
    for r, (gubun, exp, ref) in enumerate(ROWS, start=1):
        cells[(r, 0)] = dict(
            text=gubun, font_name=FONT_SB, font_size=9.5, font_color=INK,
            bold=False, align="l", anchor="ctr", **body_border,
        )
        cells[(r, 1)] = dict(
            text=exp, font_name=FONT_R, font_size=9.5, font_color=CHARCOAL,
            bold=False, align="l", anchor="ctr", **body_border,
        )
        cells[(r, 2)] = dict(
            segments=ref, font_name=FONT_R, font_size=9.5, font_color=CHARCOAL,
            bold=False, align="l", anchor="ctr", **body_border,
        )

    _shape, table = add_grid_table(
        slide, x0, table_top, full_w, table_h,
        nrows=1 + len(ROWS), ncols=3,
        cells=cells,
        col_widths=col_ratios,
        row_heights=[header_h] + row_hs,
        font_name=FONT_R, default_font_size=9.5,
        default_fill=CANVAS,
        gridlines=False,  # 수직선 없음 — 수평 hairline만 셀별 명시
    )
    # 표 상단 마감 hairline (헤더 위)
    for c in range(3):
        set_cell_border(table.cell(0, c), "t", color=FOG, width_pt=0.75)
    # 셀 행간 = 폰트 × 1.3 (estimator 라인 높이와 동일 — 넘침 없음)
    for r in range(1 + len(ROWS)):
        size = 10.5 if r == 0 else 9.5
        for c in range(3):
            for para in table.cell(r, c).text_frame.paragraphs:
                para.line_spacing = Pt(round(size * 1.3, 1))

    # ---- 3. 하단 다크 슬랩 (INK 배경 + ON_INK, '지원 가능' PRIMARY_BRIGHT) ----
    slab = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0, slab_top, full_w, slab_h)
    slab.fill.solid()
    slab.fill.fore_color.rgb = INK
    slab.line.fill.background()
    set_body_anchor(slab, "ctr")
    set_text_inset(slab)
    stf = slab.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.alignment = PP_ALIGN.CENTER
    run = sp.add_run()
    run.text = CLOSING_TEXT_HEAD
    run.font.name = FONT_R
    run.font.size = Pt(11)
    run.font.bold = False
    run.font.color.rgb = ON_INK
    run = sp.add_run()
    run.text = CLOSING_TEXT_EMPH
    run.font.name = FONT_SB
    run.font.size = Pt(11)
    run.font.bold = False
    run.font.color.rgb = PRIMARY_BRIGHT

    # ---- 마무리: EA 폰트 정리 후 맑은 고딕 일괄 강제 (라틴+EA+CS, SemiBold→bold 변환) ----
    _finalize_korean_typography(slide)
    force_font(slide, "맑은 고딕")
